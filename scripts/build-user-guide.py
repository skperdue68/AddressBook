"""Build the Word guide from docs/user-admin-guide.md.
Requires python-docx for authoring only; not used by bot startup.
Render and inspect the resulting document before committing it.
"""
from pathlib import Path
import argparse
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT = Path(__file__).resolve().parents[1]

def node(name, **attributes):
    element = OxmlElement('w:'+name)
    for key, value in attributes.items(): element.set(qn('w:'+key), str(value))
    return element

def rich(paragraph, text, color=None):
    for part in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^\)]+\))', text):
        match = re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)', part)
        if match:
            link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
            run = node('r'); props = node('rPr'); props.append(node('color', val=color or '245D91')); props.append(node('u', val='single')); run.append(props)
            content = node('t'); content.text = match[1]; run.append(content); link.append(run); paragraph._p.append(link)
        else:
            bold = part.startswith('**') and part.endswith('**'); run = paragraph.add_run(part[2:-2] if bold else part); run.bold = bold
            if color: run.font.color.rgb = RGBColor.from_string(color)

def fresh_list(doc):
    root = doc.part.numbering_part.element
    abstract_id = max([int(x.get(qn('w:abstractNumId'))) for x in root.findall(qn('w:abstractNum'))]+[-1])+1
    num_id = max([int(x.get(qn('w:numId'))) for x in root.findall(qn('w:num'))]+[0])+1
    abstract = node('abstractNum', abstractNumId=abstract_id); level = node('lvl', ilvl=0)
    for name, value in [('start',1),('numFmt','decimal'),('lvlText','%1.'),('suff','tab')]: level.append(node(name,val=value))
    props = node('pPr'); props.append(node('ind',left=360,hanging=360)); level.append(props); abstract.append(level); root.append(abstract)
    number = node('num',numId=num_id); number.append(node('abstractNumId',val=abstract_id)); root.append(number)
    return num_id

def add_table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?',cell) for cell in row)]
    table = doc.add_table(rows=0,cols=2); table.autofit = False
    table.columns[0].width = Inches(1.35); table.columns[1].width = Inches(5.55)
    borders = node('tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'): borders.append(node(edge,val='single',sz=4,color='D9D9D9'))
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row(); row._tr.get_or_add_trPr().append(node('cantSplit'))
        if not index: row._tr.get_or_add_trPr().append(node('tblHeader'))
        for col,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width = Inches(1.35 if col==0 else 5.55); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr(); props.append(node('shd',fill='17324D' if not index else ('F2F5FA' if index%2 else 'FFFFFF')))
            padding = node('tcMar')
            for edge in ('top','left','bottom','right'): padding.append(node(edge,w=100,type='dxa'))
            props.append(padding); p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2); rich(p,value,'FFFFFF' if not index else None)
            if not index:
                for run in p.runs: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def build(output):
    doc = Document(); section = doc.sections[0]
    section.page_width=Inches(8.5); section.page_height=Inches(11); section.top_margin=section.bottom_margin=Inches(.7); section.left_margin=section.right_margin=Inches(.8); section.header_distance=section.footer_distance=Inches(.3)
    normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string('172638'); normal.paragraph_format.line_spacing=1.08; normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',26),('Heading 1',18),('Heading 2',12)]:
        style=doc.styles[name]; style.font.name='Calibri'; style.font.size=Pt(size); style.font.bold=name!='Title'; style.font.color.rgb=RGBColor.from_string('000000' if name=='Title' else '17324D'); style.font.underline=False; style.paragraph_format.space_before=Pt(0 if name=='Title' else 10); style.paragraph_format.space_after=Pt(10 if name=='Title' else 5); style.paragraph_format.keep_with_next=True
    for border in list(doc.styles.element.xpath('//w:pBdr')): border.getparent().remove(border)
    header=section.header.paragraphs[0]; header.text='HEARTHBOUND   /   USER AND ADMINISTRATOR GUIDE'
    for run in header.runs: run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string('65788D')
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run('HearthBound   |   '); run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string('65788D'); footer._p.append(node('fldSimple',instr='PAGE'))
    lines=(ROOT/'docs/user-admin-guide.md').read_text(encoding='utf-8').splitlines(); i=0; code=False; numbered=False; list_id=None
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if line=='```bash': code=True; numbered=False; continue
        if line=='```': code=False; continue
        if code:
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1; p.paragraph_format.left_indent=Inches(.12); p.paragraph_format.keep_with_next=i<len(lines) and lines[i].strip()!='```'
            if not p.paragraph_format.keep_with_next: p.paragraph_format.space_after=Pt(8)
            run=p.add_run(line); run.font.name='Consolas'; run.font.size=Pt(9.5); continue
        if not line: continue
        if line=='<!-- page -->': doc.add_page_break(); numbered=False; continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'): block.append(lines[i]); i+=1
            add_table(doc,block); numbered=False; continue
        heading=next(((prefix,style) for prefix,style in [('# ','Title'),('## ','Heading 1'),('### ','Heading 2')] if line.startswith(prefix)),None)
        if heading: doc.add_paragraph(line[len(heading[0]):],heading[1]); numbered=False; continue
        match=re.match(r'^\d+\.\s+(.*)',line)
        if match:
            if not numbered: list_id=fresh_list(doc)
            p=doc.add_paragraph(); props=node('numPr'); props.append(node('ilvl',val=0)); props.append(node('numId',val=list_id)); p._p.get_or_add_pPr().append(props); rich(p,match[1]); numbered=True; continue
        numbered=False; rich(doc.add_paragraph(style='List Bullet' if line.startswith('- ') else None),line[2:] if line.startswith('- ') else line)
    doc.core_properties.title='HearthBound User and Administrator Guide'; doc.core_properties.author='HearthBound'; doc.core_properties.subject='Discord and web workflows'; output.parent.mkdir(parents=True,exist_ok=True); doc.save(output)

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__); parser.add_argument('--output',type=Path,default=ROOT/'HearthBound_User_Admin_Guide.docx'); build(parser.parse_args().output)
